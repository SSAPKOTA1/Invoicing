"""Exported numbers must equal the database / ledger."""

from __future__ import annotations

import csv
from datetime import date

import pytest

from supplier_app.reports.builders import ReportService
from supplier_app.reports.exporters import export_csv, export_pdf, export_xlsx
from supplier_app.services.demo_data import DemoDataService
from supplier_app.util.money import format_cents

TODAY = date(2026, 6, 15)


@pytest.fixture()
def demo(svc):
    DemoDataService(svc).load(TODAY)
    return svc, ReportService(svc)


def test_open_items_total_equals_ledger_and_aging(demo) -> None:
    svc, rs = demo
    rep = rs.open_items(TODAY, None)
    ledger_total = sum(s.balance.balance_cents for s in svc.ledger.invoice_states(TODAY) if s.balance.balance_cents > 0)
    assert rep.totals[9] == ledger_total == 3_277_340
    assert sum(r[9] for r in rep.raw_rows) == ledger_total
    aging = svc.ledger.aging(TODAY)
    assert [r[1] for r in rep.extra[0].raw_rows] == list(aging.values()) and rep.extra[0].totals[1] == ledger_total
    for r in rep.raw_rows:  # each line: gross + charges - settled == open
        assert r[6] + r[7] - r[8] == r[9]


def test_statement_matches_supplier_balance(demo) -> None:
    svc, rs = demo
    for p in svc.suppliers.suppliers():
        rep = rs.supplier_statement(p.id, None, None)
        bal = svc.ledger.supplier_balance(p.id).total_cents
        assert rep.totals[6] == bal
        assert rep.totals[4] - rep.totals[5] == bal
        assert rep.raw_rows[-1][6] == bal
    part = rs.supplier_statement(1, date(2026, 1, 1), None)
    assert part.raw_rows[0][3] == "Saldovortrag"
    with pytest.raises(ValueError):
        rs.build("statement")
    with pytest.raises(ValueError):
        rs.build("bogus")


def test_dunning_report_matches_reconciliation(demo) -> None:
    svc, rs = demo
    rep = rs.dunning_overview(None, None, None)
    assert len(rep.raw_rows) == len(svc.repos.notices.list()) == 11
    total_diff = sum(svc.ledger.reconcile_notice(n.id).difference_cents for n in svc.repos.notices.list())
    assert rep.totals[7] == total_diff == 200_000
    flagged = [r for r in rep.raw_rows if r[7] != 0]
    assert len(flagged) == 1 and "Zahlung" in flagged[0][12]
    assert rs.dunning_overview(date(2099, 1, 1), None, None).raw_rows == []


def test_payments_report_matches_payment_table(demo) -> None:
    svc, rs = demo
    rep = rs.payments(None, None, None)
    assert rep.totals[2] == sum(p.amount_cents for p in svc.repos.payments.list()) == 3_409_300
    assert rep.totals[5] == 40_000  # Alpha overpayment stays as credit
    assert {r[6] for r in rep.raw_rows} == {"gebucht", "Guthaben"}
    pay = svc.repos.payments.list()[0]
    svc.payments.reverse_payment(pay.id)
    rep2 = rs.payments(None, None, None)
    assert "storniert" in {r[6] for r in rep2.raw_rows}
    assert rep2.totals[2] == rep.totals[2] - pay.amount_cents


@pytest.mark.parametrize("kind", ["open_items", "dunning", "payments", "statement"])
def test_csv_xlsx_pdf_roundtrip_numbers(demo, tmp_path, kind) -> None:
    svc, rs = demo
    rep = rs.build(kind, supplier_id=1, as_of=TODAY)
    csv_path = export_csv(rep, tmp_path / "r.csv")
    with csv_path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh, delimiter=";"))
    assert rows[0] == [rep.title] and rows[2] == rep.columns
    body = rows[3:3 + len(rep.rows)]
    for exported, raw in zip(body, rep.raw_rows, strict=True):
        for i in rep.money_columns:
            if isinstance(raw[i], int):
                assert exported[i] == format_cents(raw[i], symbol=False)
    totals = rows[3 + len(rep.rows)]
    for i in rep.money_columns:
        if rep.totals[i] is not None:
            assert totals[i] == format_cents(rep.totals[i], symbol=False)

    from openpyxl import load_workbook
    wb = load_workbook(export_xlsx(rep, tmp_path / "r.xlsx"))
    ws = wb.worksheets[0]
    assert [c.value for c in ws[4]][: len(rep.columns)] == rep.columns
    for r_idx, raw in enumerate(rep.raw_rows, start=5):
        for i in rep.money_columns:
            if isinstance(raw[i], int):
                assert round(float(ws.cell(row=r_idx, column=i + 1).value) * 100) == raw[i]
    last = 5 + len(rep.raw_rows)
    for i in rep.money_columns:
        if rep.totals[i] is not None:
            assert round(float(ws.cell(row=last, column=i + 1).value) * 100) == rep.totals[i]
    assert len(wb.worksheets) == 1 + len(rep.extra)

    import pymupdf
    pdf = export_pdf(rep, tmp_path / "r.pdf")
    with pymupdf.open(pdf) as doc:
        text = "\n".join(page.get_text() for page in doc)
    assert rep.title in text
    for i in rep.money_columns:
        if rep.totals[i] is not None and isinstance(rep.totals[i], int):
            assert format_cents(rep.totals[i]) in text
